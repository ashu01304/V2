from __future__ import annotations

import argparse
import math
import sys
from collections import defaultdict, deque
from dataclasses import dataclass
from pathlib import Path
from typing import Iterable

import pandas as pd
from openpyxl import load_workbook
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo


REQUIRED = {
    "transacttime": "transactTime",
    "exchangedisplay": "exchangeDisplay",
    "contract": "contract",
    "side": "side",
    "displayfilledqty": "displayFilledQty",
    "displayprice": "displayPrice",
    "execqty": "execQty",
}


@dataclass
class Lot:
    time: pd.Timestamp
    side: str
    qty: float
    price: float
    source_row: int


def normalized_name(value: object) -> str:
    return "".join(ch for ch in str(value).strip().lower() if ch.isalnum())


def align_required_columns(frame: pd.DataFrame) -> pd.DataFrame | None:
    """Return a copy with canonical column names, or None when columns are missing."""
    lookup = {normalized_name(col): col for col in frame.columns}
    if not set(REQUIRED).issubset(lookup):
        return None
    rename = {lookup[key]: canonical for key, canonical in REQUIRED.items()}
    return frame.rename(columns=rename)


def load_source(path: Path) -> tuple[str, pd.DataFrame]:
    """Load a CSV/TSV file or find the first valid worksheet in an Excel file."""
    suffix = path.suffix.lower()
    if suffix in {".csv", ".txt", ".tsv"}:
        separators = ["\t"] if suffix == ".tsv" else [None, ",", ";", "\t", "|"]
        errors: list[str] = []
        for encoding in ("utf-8-sig", "utf-8", "cp1252"):
            for separator in separators:
                try:
                    frame = pd.read_csv(path, sep=separator, engine="python", encoding=encoding)
                    aligned = align_required_columns(frame)
                    if aligned is not None:
                        return f"CSV ({encoding})", aligned
                    errors.append(f"{encoding}, separator={separator!r}: {list(frame.columns)}")
                except (UnicodeDecodeError, pd.errors.ParserError) as exc:
                    errors.append(f"{encoding}, separator={separator!r}: {exc}")
        raise ValueError(
            "CSV does not contain all required columns. Expected: "
            + ", ".join(REQUIRED.values())
            + "\nAttempts:\n"
            + "\n".join(errors[:12])
        )

    if suffix not in {".xlsx", ".xlsm", ".xls"}:
        raise ValueError("Input must be a .csv, .tsv, .txt, .xlsx, .xlsm, or .xls file")

    workbook = pd.ExcelFile(path)
    inspected: list[str] = []
    for sheet in workbook.sheet_names:
        frame = pd.read_excel(path, sheet_name=sheet)
        inspected.append(f"{sheet}: {', '.join(map(str, frame.columns))}")
        aligned = align_required_columns(frame)
        if aligned is not None:
            return sheet, aligned
    details = "\n".join(inspected)
    raise ValueError(
        "No worksheet contains all required columns. Expected: "
        + ", ".join(REQUIRED.values())
        + f"\nSheets inspected:\n{details}"
    )


def clean_executions(frame: pd.DataFrame) -> pd.DataFrame:
    data = frame[list(REQUIRED.values())].copy()
    data["Source Row"] = data.index + 2
    data["transactTime"] = pd.to_datetime(data["transactTime"], errors="coerce")
    data["exchangeDisplay"] = data["exchangeDisplay"].astype("string").str.strip()
    data["contract"] = data["contract"].astype("string").str.strip()
    data["side"] = data["side"].astype("string").str.strip().str.title()
    data["displayPrice"] = pd.to_numeric(data["displayPrice"], errors="coerce")
    data["execQty"] = pd.to_numeric(data["execQty"], errors="coerce")
    data["displayFilledQty"] = pd.to_numeric(data["displayFilledQty"], errors="coerce")
    data["Quantity"] = data["execQty"].where(data["execQty"] > 0, data["displayFilledQty"])

    invalid = data[
        data["transactTime"].isna()
        | data["exchangeDisplay"].isna()
        | data["contract"].isna()
        | ~data["side"].isin(["Buy", "Sell"])
        | data["Quantity"].isna()
        | (data["Quantity"] <= 0)
        | data["displayPrice"].isna()
    ]
    if not invalid.empty:
        rows = ", ".join(map(str, invalid["Source Row"].astype(int).tolist()[:20]))
        raise ValueError(f"Invalid or incomplete execution data on source row(s): {rows}")

    data["Signed Qty"] = data["Quantity"].where(data["side"].eq("Buy"), -data["Quantity"])
    data["Transaction Value"] = data["displayPrice"] * data["Signed Qty"]
    data = data.sort_values(["transactTime", "Source Row"], kind="stable").reset_index(drop=True)
    return data


def fifo_match(data: pd.DataFrame) -> tuple[pd.DataFrame, pd.DataFrame]:
    queues: dict[tuple[str, str], deque[Lot]] = defaultdict(deque)
    matches: list[dict] = []

    for row in data.to_dict("records"):
        exchange = row["exchangeDisplay"]
        contract = row["contract"]
        side = row["side"]
        quantity = float(row["Quantity"])
        price = float(row["displayPrice"])
        time = row["transactTime"]
        source_row = int(row["Source Row"])
        queue = queues[(exchange, contract)]

        while quantity > 1e-12 and queue and queue[0].side != side:
            opening = queue[0]
            closed_qty = min(quantity, opening.qty)
            position = "Long" if opening.side == "Buy" else "Short"
            pnl = ((price - opening.price) if position == "Long" else (opening.price - price)) * closed_qty
            matches.append(
                {
                    "Exchange": exchange,
                    "Contract": contract,
                    "Position": position,
                    "Quantity": closed_qty,
                    "Entry Time": opening.time,
                    "Entry Side": opening.side,
                    "Entry Price": opening.price,
                    "Exit Time": time,
                    "Exit Side": side,
                    "Exit Price": price,
                    "Realized Price P&L": pnl,
                    "Entry Source Row": opening.source_row,
                    "Exit Source Row": source_row,
                }
            )
            opening.qty -= closed_qty
            quantity -= closed_qty
            if opening.qty <= 1e-12:
                queue.popleft()

        if quantity > 1e-12:
            queue.append(Lot(time=time, side=side, qty=quantity, price=price, source_row=source_row))

    open_rows: list[dict] = []
    for (exchange, contract), queue in queues.items():
        for lot in queue:
            open_rows.append(
                {
                    "Exchange": exchange,
                    "Contract": contract,
                    "Open Side": "Long" if lot.side == "Buy" else "Short",
                    "Quantity": lot.qty,
                    "Entry Time": lot.time,
                    "Entry Price": lot.price,
                    "Source Row": lot.source_row,
                }
            )

    matched = pd.DataFrame(matches)
    open_lots = pd.DataFrame(open_rows)
    if not matched.empty:
        matched.insert(0, "Trade No.", range(1, len(matched) + 1))
    if not open_lots.empty:
        open_lots = open_lots.sort_values(["Exchange", "Contract", "Entry Time"]).reset_index(drop=True)
        open_lots.insert(0, "Lot No.", range(1, len(open_lots) + 1))
    return matched, open_lots


def weighted_average(group: pd.DataFrame, side: str) -> float:
    subset = group[group["side"].eq(side)]
    if subset.empty:
        return math.nan
    return float((subset["displayPrice"] * subset["Quantity"]).sum() / subset["Quantity"].sum())


def build_summary(data: pd.DataFrame, matched: pd.DataFrame, open_lots: pd.DataFrame) -> pd.DataFrame:
    realized_lookup = (
        matched.groupby(["Exchange", "Contract"])["Realized Price P&L"].sum().to_dict()
        if not matched.empty else {}
    )
    open_avg: dict[tuple[str, str], float] = {}
    if not open_lots.empty:
        for key, group in open_lots.groupby(["Exchange", "Contract"], sort=False):
            open_avg[key] = float((group["Entry Price"] * group["Quantity"]).sum() / group["Quantity"].sum())

    rows: list[dict] = []
    for key, group in data.groupby(["exchangeDisplay", "contract"], sort=True):
        exchange, contract = key
        buy_qty = float(group.loc[group["side"].eq("Buy"), "Quantity"].sum())
        sell_qty = float(group.loc[group["side"].eq("Sell"), "Quantity"].sum())
        net_qty = buy_qty - sell_qty
        rows.append(
            {
                "Exchange": exchange,
                "Contract": contract,
                "Buy Qty": buy_qty,
                "Avg Buy Price": weighted_average(group, "Buy"),
                "Sell Qty": sell_qty,
                "Avg Sell Price": weighted_average(group, "Sell"),
                "Net Qty": net_qty,
                "Net Side": "Long" if net_qty > 1e-12 else "Short" if net_qty < -1e-12 else "Flat",
                "Open Avg Price": open_avg.get(key, math.nan),
                "Last Price": float(group.iloc[-1]["displayPrice"]),
                "Realized Price P&L": realized_lookup.get(key, 0.0),
                "Executions": int(len(group)),
                "First Transaction": group["transactTime"].min(),
                "Last Transaction": group["transactTime"].max(),
            }
        )
    return pd.DataFrame(rows)


def add_excel_table(ws, name: str) -> None:
    if ws.max_row < 2 or ws.max_column < 1:
        return
    ref = f"A1:{get_column_letter(ws.max_column)}{ws.max_row}"
    table = Table(displayName=name, ref=ref)
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2", showFirstColumn=False, showLastColumn=False,
        showRowStripes=True, showColumnStripes=False
    )
    ws.add_table(table)


def format_workbook(path: Path) -> None:
    wb = load_workbook(path)
    widths = {
        "Contract": 52, "Entry Time": 23, "Exit Time": 23,
        "First Transaction": 23, "Last Transaction": 23,
        "transactTime": 23, "Exchange": 12, "exchangeDisplay": 16,
        "Net Side": 12, "Open Side": 12,
    }
    for index, ws in enumerate(wb.worksheets, start=1):
        ws.freeze_panes = "A2"
        ws.sheet_view.showGridLines = False
        ws.auto_filter.ref = ws.dimensions
        for cell in ws[1]:
            cell.fill = PatternFill("solid", fgColor="1F4E78")
            cell.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.row_dimensions[1].height = 30
        for col in range(1, ws.max_column + 1):
            header = str(ws.cell(1, col).value or "")
            sample_lengths = [len(str(ws.cell(row, col).value or "")) for row in range(1, min(ws.max_row, 80) + 1)]
            ws.column_dimensions[get_column_letter(col)].width = widths.get(header, min(max(sample_lengths + [10]) + 2, 24))
            if "Time" in header or header == "transactTime":
                for row in range(2, ws.max_row + 1):
                    ws.cell(row, col).number_format = "dd-mmm-yy hh:mm:ss.000"
            if "Price" in header or "P&L" in header or header == "Transaction Value":
                for row in range(2, ws.max_row + 1):
                    ws.cell(row, col).number_format = "0.0000;[Red](0.0000);-"
        add_excel_table(ws, f"Table{index}")

    summary = wb["Net Positions"]
    headers = {cell.value: cell.column for cell in summary[1]}
    if "Net Side" in headers:
        col = get_column_letter(headers["Net Side"])
        rng = f"{col}2:{col}{summary.max_row}"
        summary.conditional_formatting.add(rng, FormulaRule(formula=[f'{col}2="Long"'], fill=PatternFill("solid", fgColor="E2F0D9")))
        summary.conditional_formatting.add(rng, FormulaRule(formula=[f'{col}2="Short"'], fill=PatternFill("solid", fgColor="FCE4D6")))
    if "Realized Price P&L" in headers:
        col = get_column_letter(headers["Realized Price P&L"])
        rng = f"{col}2:{col}{summary.max_row}"
        summary.conditional_formatting.add(rng, CellIsRule(operator="greaterThan", formula=["0"], font=Font(color="008000", bold=True)))
        summary.conditional_formatting.add(rng, CellIsRule(operator="lessThan", formula=["0"], font=Font(color="C00000", bold=True)))
    wb.save(path)


def write_report(output: Path, data: pd.DataFrame, summary: pd.DataFrame, matched: pd.DataFrame, open_lots: pd.DataFrame) -> None:
    normalized = data.rename(columns={
        "transactTime": "Transaction Time", "exchangeDisplay": "Exchange",
        "contract": "Contract", "side": "Side", "displayFilledQty": "Display Filled Qty",
        "displayPrice": "Price", "execQty": "Exec Qty",
    })
    normalized = normalized[[
        "Source Row", "Transaction Time", "Exchange", "Contract", "Side",
        "Display Filled Qty", "Exec Qty", "Quantity", "Price", "Signed Qty", "Transaction Value"
    ]]
    with pd.ExcelWriter(output, engine="openpyxl", datetime_format="dd-mmm-yy hh:mm:ss.000") as writer:
        summary.to_excel(writer, sheet_name="Net Positions", index=False)
        (matched if not matched.empty else pd.DataFrame(columns=[
            "Trade No.", "Exchange", "Contract", "Position", "Quantity", "Entry Time",
            "Entry Side", "Entry Price", "Exit Time", "Exit Side", "Exit Price",
            "Realized Price P&L", "Entry Source Row", "Exit Source Row"
        ])).to_excel(writer, sheet_name="Matched Trades", index=False)
        (open_lots if not open_lots.empty else pd.DataFrame(columns=[
            "Lot No.", "Exchange", "Contract", "Open Side", "Quantity", "Entry Time", "Entry Price", "Source Row"
        ])).to_excel(writer, sheet_name="Open Lots", index=False)
        normalized.to_excel(writer, sheet_name="Normalized Executions", index=False)
    format_workbook(output)


def main(argv: Iterable[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="Net every contract in a broker execution CSV or Excel file.")
    parser.add_argument("input", type=Path, help="Broker execution .csv, .tsv, .txt, or Excel file")
    parser.add_argument("-o", "--output", type=Path, help="Output .xlsx file")
    args = parser.parse_args(argv)
    if not args.input.exists():
        parser.error(f"Input file not found: {args.input}")
    output = args.output or args.input.with_name(f"{args.input.stem}_net_transactions.xlsx")
    try:
        source_name, source = load_source(args.input)
        executions = clean_executions(source)
        matched, open_lots = fifo_match(executions)
        summary = build_summary(executions, matched, open_lots)
        write_report(output, executions, summary, matched, open_lots)
    except Exception as exc:
        print(f"Error: {exc}", file=sys.stderr)
        return 1
    print(f"Source: {source_name}")
    print(f"Executions: {len(executions):,}")
    print(f"Contracts: {len(summary):,}")
    print(f"Matched trades: {len(matched):,}")
    print(f"Open lots: {len(open_lots):,}")
    print(f"Report: {output.resolve()}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
